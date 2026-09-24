import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

from app.agents.jd_agent import analyze_jd  # noqa: E402
from app.config import override  # noqa: E402
from app.pipeline import run  # noqa: E402
from app.reporting.export import (  # noqa: E402
    RANKING_COLUMNS,
    REQUIREMENT_COLUMNS,
    build_ranking_rows,
    build_requirement_rows,
    to_csv,
    to_excel,
)

override(provider="mock")

JD = ROOT / "data" / "jd" / "senior_ai_engineer.md"
RESUMES = sorted(str(p) for p in (ROOT / "data" / "resumes").glob("*.pdf"))


def _outcome(n=3):
    return run(str(JD), RESUMES[:n], enrich=False)


# ------------------------------------------------------- heading filtering

LINKEDIN_STYLE = """AI Engineer
About The Role
We are hiring an AI Engineer.
What You'll Do
- Design and deploy machine learning models for business decisions
- Build recommendation engines and classification models
Cross-Functional Collaboration
- Partner with product and engineering teams
Requirements
- 3+ years of experience in machine learning
- Strong Python and SQL
Benefits
- Health insurance and equity
"""


def test_section_headings_are_not_treated_as_requirements():
    job = analyze_jd(LINKEDIN_STYLE)
    skills = " | ".join(r.skill.lower() for r in job.requirements)
    for heading in ("about the role", "what you'll do", "cross-functional collaboration",
                    "benefits", "requirements"):
        assert heading not in skills, f"heading leaked into requirements: {heading}"


def test_real_requirements_still_extracted():
    job = analyze_jd(LINKEDIN_STYLE)
    skills = " ".join(r.skill.lower() for r in job.requirements)
    assert "python" in skills
    assert "machine learning" in skills
    assert job.minimum_experience_years == 3.0


# -------------------------------------------------------------- export rows

def test_ranking_rows_have_every_column():
    outcome = _outcome()
    rows = build_ranking_rows(outcome.job, outcome.ranked)
    assert len(rows) == len(outcome.ranked)
    for row in rows:
        for col in RANKING_COLUMNS:
            assert col in row, f"missing column {col}"


def test_experience_comparison_is_explicit():
    """The JD asks for 5+ years; the table must show required vs actual."""
    outcome = _outcome()
    rows = build_ranking_rows(outcome.job, outcome.ranked)
    row = rows[0]
    assert row["Required experience (yrs)"] == 5.0
    assert row["Experience met"] in ("Yes", "No", "Could not determine")


def test_requirement_rows_include_an_experience_row_per_candidate():
    outcome = _outcome()
    rows = build_requirement_rows(outcome.job, outcome.ranked)
    exp_rows = [r for r in rows if r["Requirement"] == "Total professional experience"]
    assert len(exp_rows) == len(outcome.ranked)


def test_every_requirement_row_states_satisfied():
    outcome = _outcome()
    rows = build_requirement_rows(outcome.job, outcome.ranked)
    assert rows
    for row in rows:
        assert row["Satisfied"] in (
            "Yes", "Partial", "Weak", "No", "Not specified in JD", "Could not determine"
        )
        for col in REQUIREMENT_COLUMNS:
            assert col in row


def test_ranking_is_ordered_by_rank():
    outcome = run(str(JD), RESUMES, enrich=False)
    rows = build_ranking_rows(outcome.job, outcome.ranked)
    assert [r["Rank"] for r in rows] == list(range(1, len(rows) + 1))
    scores = [r["Score /100"] for r in rows]
    assert scores == sorted(scores, reverse=True)


# ------------------------------------------------------------------ writers

def test_csv_export_is_well_formed():
    outcome = _outcome()
    text = to_csv(build_requirement_rows(outcome.job, outcome.ranked), REQUIREMENT_COLUMNS)
    lines = text.splitlines()
    assert lines[0].startswith("Rank,Candidate,Requirement")
    assert len(lines) > 5


def test_excel_workbook_has_three_sheets():
    import io

    from openpyxl import load_workbook

    outcome = _outcome()
    data = to_excel(outcome.job, outcome.ranked)
    assert data and len(data) > 5000
    wb = load_workbook(io.BytesIO(data))
    assert wb.sheetnames == ["Ranking", "Requirements", "Job"]
    assert wb["Requirements"].max_row > 5
    assert wb["Ranking"].freeze_panes == "A2"
