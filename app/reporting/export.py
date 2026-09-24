"""Tabular export of an evaluation run.

Three sheets, because recruiters read them in this order:

  Ranking       one row per candidate — score, verdict, how many requirements
                are actually evidenced, experience against the JD's minimum
  Requirements  one row per candidate per requirement — what was asked, what
                the candidate showed, whether it is satisfied, and the quote
  Job           what the system read out of the job description

No formulas: this is a record of a completed run, not a model to recalculate.
Values are written as values so the file opens identically in Excel, Numbers,
Google Sheets and pandas.
"""

from __future__ import annotations

import csv
import io
from datetime import datetime
from typing import Any, Dict, List, Optional

from app.schemas.evidence import EvidenceStrength
from app.schemas.job import JobDescription

# How an evidence level reads in a spreadsheet cell.
SATISFIED: Dict[EvidenceStrength, str] = {
    EvidenceStrength.STRONG_PRODUCTION_EVIDENCE: "Yes",
    EvidenceStrength.PROFESSIONAL_EVIDENCE: "Yes",
    EvidenceStrength.PROJECT_EVIDENCE: "Partial",
    EvidenceStrength.MENTIONED: "Weak",
    EvidenceStrength.NOT_FOUND: "No",
}

EVIDENCE_LABEL: Dict[EvidenceStrength, str] = {
    EvidenceStrength.STRONG_PRODUCTION_EVIDENCE: "Shipped in production",
    EvidenceStrength.PROFESSIONAL_EVIDENCE: "Used professionally",
    EvidenceStrength.PROJECT_EVIDENCE: "Project work",
    EvidenceStrength.MENTIONED: "Listed only, no described work",
    EvidenceStrength.NOT_FOUND: "Not found in resume or links",
}

RANKING_COLUMNS = [
    "Rank", "Candidate", "Score /100", "Verdict",
    "Required areas evidenced", "Required areas total", "Coverage %",
    "Required experience (yrs)", "Candidate experience (yrs)", "Experience met",
    "Must-have points", "Nice-to-have points", "Experience points",
    "External evidence points", "Flags", "GitHub", "Portfolio", "LinkedIn",
]

REQUIREMENT_COLUMNS = [
    "Rank", "Candidate", "Requirement", "Type", "Weight",
    "Evidence level", "Satisfied", "Evidence found", "Source",
]


def _experience_row(job: JobDescription, result) -> Dict[str, Any]:
    required = job.minimum_experience_years
    actual = result.profile.total_experience_years if result.profile else None
    if required is None:
        met = "Not specified in JD"
    elif actual is None:
        met = "Could not determine"
    else:
        met = "Yes" if actual >= required else "No"
    return {"required": required, "actual": actual, "met": met}


def build_ranking_rows(job: JobDescription, ranked: List) -> List[Dict[str, Any]]:
    rows = []
    for i, result in enumerate(ranked, start=1):
        report = result.report
        v = result.verdict
        exp = _experience_row(job, result)
        links = result.links()
        total = (v.must_have_total if v else len(job.must_haves)) or 1
        hit = v.must_have_hit if v else 0
        rows.append(
            {
                "Rank": i,
                "Candidate": result.display_name,
                "Score /100": report.score,
                "Verdict": v.label if v else "",
                "Required areas evidenced": hit,
                "Required areas total": total,
                "Coverage %": round(hit / total * 100, 1),
                "Required experience (yrs)": exp["required"] if exp["required"] else "Not specified",
                "Candidate experience (yrs)": exp["actual"] if exp["actual"] is not None else "Unknown",
                "Experience met": exp["met"],
                "Must-have points": report.breakdown.must_have_coverage,
                "Nice-to-have points": report.breakdown.nice_to_have_coverage,
                "Experience points": report.breakdown.experience_fit,
                "External evidence points": report.breakdown.external_evidence,
                "Flags": len(report.flags),
                "GitHub": links.get("github", "—"),
                "Portfolio": links.get("portfolio", "—"),
                "LinkedIn": links.get("linkedin", "—"),
            }
        )
    return rows


def build_requirement_rows(job: JobDescription, ranked: List) -> List[Dict[str, Any]]:
    rows = []
    for i, result in enumerate(ranked, start=1):
        report = result.report

        # The experience gate is a requirement like any other, so it gets a row.
        exp = _experience_row(job, result)
        rows.append(
            {
                "Rank": i,
                "Candidate": result.display_name,
                "Requirement": "Total professional experience",
                "Type": "Must have" if exp["required"] else "Informational",
                "Weight": "—",
                "Evidence level": (
                    f"{exp['actual']} years" if exp["actual"] is not None else "Unknown"
                ),
                "Satisfied": exp["met"],
                "Evidence found": (
                    f"JD asks for {exp['required']}+ years; resume shows {exp['actual']}."
                    if exp["required"] and exp["actual"] is not None
                    else "No minimum stated in the JD."
                    if not exp["required"]
                    else "Could not determine total experience from the resume."
                ),
                "Source": "resume",
            }
        )

        for req in job.must_haves + job.nice_to_haves:
            match = report.match(req.id)
            if not match:
                continue
            quote = match.evidence[0].quote if match.evidence else ""
            source = match.evidence[0].source.value if match.evidence else "—"
            rows.append(
                {
                    "Rank": i,
                    "Candidate": result.display_name,
                    "Requirement": req.skill,
                    "Type": "Must have" if req.kind.value == "must_have" else "Nice to have",
                    "Weight": req.weight,
                    "Evidence level": EVIDENCE_LABEL[match.strength],
                    "Satisfied": SATISFIED[match.strength],
                    "Evidence found": quote or "—",
                    "Source": source,
                }
            )
    return rows


def build_job_rows(job: JobDescription) -> List[Dict[str, Any]]:
    rows = [
        {"Field": "Role", "Value": job.role_title},
        {"Field": "Minimum experience (yrs)",
         "Value": job.minimum_experience_years or "Not specified"},
        {"Field": "Must-have requirements", "Value": len(job.must_haves)},
        {"Field": "Nice-to-have requirements", "Value": len(job.nice_to_haves)},
        {"Field": "Evaluated on", "Value": datetime.now().strftime("%Y-%m-%d %H:%M")},
        {"Field": "", "Value": ""},
    ]
    for req in job.must_haves:
        rows.append({"Field": f"Must have · weight {req.weight}", "Value": req.skill})
    for req in job.nice_to_haves:
        rows.append({"Field": f"Nice to have · weight {req.weight}", "Value": req.skill})
    return rows


# ------------------------------------------------------------------- writers

def to_csv(rows: List[Dict[str, Any]], columns: Optional[List[str]] = None) -> str:
    if not rows:
        return ""
    columns = columns or list(rows[0].keys())
    buf = io.StringIO()
    writer = csv.DictWriter(buf, fieldnames=columns, extrasaction="ignore")
    writer.writeheader()
    writer.writerows(rows)
    return buf.getvalue()


def to_excel(job: JobDescription, ranked: List) -> Optional[bytes]:
    """Three-sheet workbook. Returns None if openpyxl is unavailable."""
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter
    except ImportError:
        return None

    header_font = Font(name="Arial", bold=True, color="FFFFFF", size=10)
    header_fill = PatternFill("solid", fgColor="1F3A5F")
    body_font = Font(name="Arial", size=10)
    wrap = Alignment(wrap_text=True, vertical="top")

    fills = {
        "Yes": PatternFill("solid", fgColor="D8EFDF"),
        "Partial": PatternFill("solid", fgColor="FBEFD2"),
        "Weak": PatternFill("solid", fgColor="FBE4D0"),
        "No": PatternFill("solid", fgColor="F6D9D9"),
    }

    wb = Workbook()

    def write_sheet(title, columns, rows, widths, satisfied_col=None, freeze="A2"):
        ws = wb.create_sheet(title)
        ws.append(columns)
        for cell in ws[1]:
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(vertical="center")
        for row in rows:
            ws.append([row.get(c, "") for c in columns])
        for r in ws.iter_rows(min_row=2):
            for cell in r:
                cell.font = body_font
                cell.alignment = wrap
        if satisfied_col and satisfied_col in columns:
            idx = columns.index(satisfied_col) + 1
            for r in range(2, ws.max_row + 1):
                cell = ws.cell(row=r, column=idx)
                if cell.value in fills:
                    cell.fill = fills[cell.value]
                    cell.font = Font(name="Arial", size=10, bold=True)
        for i, width in enumerate(widths, start=1):
            ws.column_dimensions[get_column_letter(i)].width = width
        ws.freeze_panes = freeze
        ws.auto_filter.ref = ws.dimensions
        return ws

    wb.remove(wb.active)

    write_sheet(
        "Ranking", RANKING_COLUMNS, build_ranking_rows(job, ranked),
        [6, 26, 11, 20, 14, 14, 11, 14, 14, 13, 14, 14, 13, 15, 7, 34, 30, 34],
        satisfied_col="Experience met",
    )
    write_sheet(
        "Requirements", REQUIREMENT_COLUMNS, build_requirement_rows(job, ranked),
        [6, 22, 40, 13, 8, 26, 11, 70, 11],
        satisfied_col="Satisfied",
    )
    write_sheet(
        "Job", ["Field", "Value"], build_job_rows(job), [32, 70], freeze="A2"
    )

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
